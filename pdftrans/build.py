# -*- coding: utf-8 -*-
"""把 IR 渲染成 DOCX（Word 原生公式）与 HTML。

沿用今天验证过的做法：
* 色块用「单格表格 + w:tblBorders + w:shd」复刻 tcolorbox；
* 相邻表格必须插 spacer 段落，否则 Word 会合并；
* 嵌套表格的 tblW 改为 pct，避免溢出父单元格；
* LaTeX → MathML → OMML（并修补 mathml2omml 的 groupChr 缺陷）。
"""
import io, os, re, base64, warnings
warnings.filterwarnings('ignore')

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement, parse_xml

import latex2mathml.converter as conv
import mathml2omml

from .extract import Doc, Block
from .translate import strip_bullet

MNS = ' xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'
_body_latin = 'Times New Roman'
_body_cjk = '宋体'
_head_cjk = '黑体'

ALIGN = {'center': WD_ALIGN_PARAGRAPH.CENTER, 'right': WD_ALIGN_PARAGRAPH.RIGHT,
         'left': WD_ALIGN_PARAGRAPH.LEFT}
_mcache = {}

def _repair(xml):
    return re.sub(r'(<m:groupChrPr>.*?)</m:groupChr>', r'\1</m:groupChrPr>', xml, flags=re.S)

def omml(latex):
    if latex not in _mcache:
        x = _repair(mathml2omml.convert(conv.convert(latex)))
        m = re.match(r'^<m:oMath>(.*)</m:oMath>$', x, re.S)
        if not m:
            raise ValueError('bad OMML')
        _mcache[latex] = m.group(1)
    return _mcache[latex]

def mml_html(latex, display=False):
    x = conv.convert(latex)
    return x.replace('display="inline"', 'display="block"' if display else 'display="inline"')

# ───────────────────────── 模板解析 ─────────────────────────
def parse_template(t, refs):
    """→ [( 'text', txt, bold, italic ) | ( 'ref', run )]"""
    out, buf, bold, italic = [], [], False, False
    i = 0
    while i < len(t):
        if t.startswith('{{', i):
            j = t.find('}}', i)
            if j < 0:
                buf.append(t[i:]); break
            if buf:
                out.append(('text', ''.join(buf), bold, italic)); buf = []
            idx = t[i + 2:j]
            r = refs.get(int(idx)) if idx.isdigit() else None
            if r is not None:
                out.append(('ref', r))
            i = j + 2
        elif t.startswith('[[b]]', i):
            if buf:
                out.append(('text', ''.join(buf), bold, italic)); buf = []
            bold = True; i += 5
        elif t.startswith('[[/b]]', i):
            if buf:
                out.append(('text', ''.join(buf), bold, italic)); buf = []
            bold = False; i += 6
        elif t.startswith('[[i]]', i):
            if buf:
                out.append(('text', ''.join(buf), bold, italic)); buf = []
            italic = True; i += 5
        elif t.startswith('[[/i]]', i):
            if buf:
                out.append(('text', ''.join(buf), bold, italic)); buf = []
            italic = False; i += 6
        else:
            buf.append(t[i]); i += 1
    if buf:
        out.append(('text', ''.join(buf), bold, italic))
    return out

# ───────────────────────── DOCX ─────────────────────────
def _split_big(blk):
    """按「超高图」把段落内容切段。

    cases/多行公式合并后会产生 2 行以上高度的行内图；挤在文字基线上时
    图会探到文字上方（看起来像公式上下有乱码/串行）。超过 ~2 倍字号的图
    提升为独立居中大图，前后的文字各自成段；紧跟大图之后的**纯标点**碎片
    （逗号、句号）并入大图段落，避免再出现一行只有标点的小块。
    → [('inline', [seg...]) | ('big', ref, [suffix seg...])]
    """
    parts, cur = [], []
    limit = 2.0 * (blk.size or 11)
    segs = parse_template(blk.template, blk.refs)
    # 项内没有真实文字（只有项目符号/标点）⇒ 不提升：图片留在符号旁，
    # 否则会先输出一个只剩符号的空列表项
    plain = ''.join(s[1] for s in segs if s[0] == 'text')
    plain = re.sub(r'\{\{\d+\}\}|[[\]/bi]', '', plain)
    plain = re.sub(r'[\s•●○▪▫■□◦‣∙⋅·\.,;:!?…\-–—]', '', plain)
    if not plain:
        return [('inline', segs)]
    for seg in segs:
        if seg[0] == 'ref' and seg[1].kind == 'image' and seg[1].h > limit:
            if cur:
                parts.append(('inline', cur)); cur = []
            parts.append(('big', seg[1], []))
        else:
            cur.append(seg)
    if cur:
        parts.append(('inline', cur))
    # 纯标点的尾随碎片并进前一个「大图」段
    out = []
    for p in parts:
        if p[0] == 'inline' and out and out[-1][0] == 'big':
            txt = ''.join(s[1] for s in p[1] if s[0] == 'text')
            if txt and not re.sub(r'[\s\.,;:!?…、，。；：·\-–—\[\]()]', '', txt):
                out[-1][2].extend(p[1])
                continue
        out.append(p)
    return out

def _set_run(run, size=None, bold=None, italic=None, color=None, cjk=None):
    f = run.font
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if color: f.color.rgb = RGBColor.from_string(color)
    f.name = _body_latin
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rPr.insert(0, rf)
    for k, v in (('w:ascii', _body_latin), ('w:hAnsi', _body_latin),
                 ('w:cs', _body_latin), ('w:eastAsia', cjk or _body_cjk)):
        rf.set(qn(k), v)

def _borders(table, **spec):
    tblPr = table._tbl.tblPr
    bd = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + side)
        v = spec.get(side)
        if v:
            e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(v[0]))
            e.set(qn('w:space'), '0'); e.set(qn('w:color'), v[1])
        else:
            e.set(qn('w:val'), 'none'); e.set(qn('w:sz'), '0')
            e.set(qn('w:space'), '0'); e.set(qn('w:color'), 'auto')
        bd.append(e)
    tblPr.append(bd)

def _props(table, fill=None, mar=(113, 113, 57, 57)):
    tblPr = table._tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    w = OxmlElement('w:tblW'); w.set(qn('w:w'), str(int(TW * 567))); w.set(qn('w:type'), 'dxa')
    tblPr.append(w)
    cm = OxmlElement('w:tblCellMar')
    for tag, val in zip(('top', 'left', 'bottom', 'right'), mar):
        e = OxmlElement('w:' + tag); e.set(qn('w:w'), str(val)); e.set(qn('w:type'), 'dxa')
        cm.append(e)
    tblPr.append(cm)
    if fill:
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
        tblPr.append(shd)

def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def _xml_safe(t):
    """剔除 XML 1.0 不允许的字符（如 PDF 文本层里偶发的 0x01 控制符）。

    python-docx 用的是 lxml，遇到这类字符会直接抛
    "All strings must be XML compatible" 让整份 docx 生成失败；
    HTML 里它们也会变成乱码方块。注意不可删的合法字符：\\t \\n \\r。
    """
    if not t:
        return t
    return ''.join(ch for ch in t
                   if ch in '\t\n\r'
                   or 0x20 <= ord(ch) <= 0xD7FF
                   or 0xE000 <= ord(ch) <= 0xFFFD
                   or ord(ch) >= 0x10000)


def _field(p, code):
    """插入可自动更新的 Word 域，如 PAGE / NUMPAGES。"""
    r = p.add_run()
    _set_run(r, size=9)
    fld = OxmlElement('w:fldChar'); fld.set(qn('w:fldCharType'), 'begin')
    r._r.append(fld)
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve')
    it.text = ' %s ' % code
    r._r.append(it)
    fld2 = OxmlElement('w:fldChar'); fld2.set(qn('w:fldCharType'), 'separate')
    r._r.append(fld2)
    t = OxmlElement('w:t'); t.text = '1'
    r._r.append(t)
    fld3 = OxmlElement('w:fldChar'); fld3.set(qn('w:fldCharType'), 'end')
    r._r.append(fld3)

def _clear(p):
    p._p.getparent().remove(p._p)

def _emit(container, blk, space_after=3):
    size = blk.size or 11
    is_list = (blk.role == 'list')
    align = ALIGN.get(blk.align, WD_ALIGN_PARAGRAPH.LEFT)
    head_done = False
    first_inline = True
    first_p = None
    for typ, *rest in _split_big(blk):
        if typ == 'big':
            # 超高图：独立居中段落，不再挤在文字基线上
            ref, suffix = rest[0], rest[1]
            ip = container.add_paragraph()
            ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ip.paragraph_format.space_before = Pt(3)
            ip.paragraph_format.space_after = Pt(3)
            try:
                ip.add_run().add_picture(io.BytesIO(ref.png), width=Pt(ref.w))
            except Exception:
                pass
            for kind, *r2 in suffix:
                if kind == 'text' and r2 and r2[0]:
                    sr = ip.add_run(_xml_safe(r2[0]))
                    _set_run(sr, size=size, bold=blk.bold, italic=blk.italic,
                             color=blk.color)
            continue
        payload = rest[0]
        p = container.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(space_after)
        pf.line_spacing = 1.15
        if is_list and first_inline:
            # 用 Word 原生项目符号，而不是把 • 当正文字符写进去
            try:
                p.style = 'List Bullet'
            except Exception:
                pf.left_indent = Cm(0.9)
                pf.first_line_indent = Cm(-0.5)
        elif is_list:
            pf.left_indent = Cm(1.2)          # 超高图之后的续行只缩进、不加符号
        p.alignment = align
        if blk.display:
            pf.space_before = Pt(6); pf.space_after = Pt(6)
        for kind, *rest in payload:
            if kind == 'text':
                txt, bold, italic = rest
                if is_list and not head_done:
                    # 项目符号由 Word 的列表样式输出，文本里的 • 不要重复出现
                    txt = strip_bullet(txt)[0]
                    head_done = True
                if not txt:
                    continue
                r = p.add_run(_xml_safe(txt))
                _set_run(r, size=size, bold=(bold or blk.bold), italic=(italic or blk.italic),
                         color=blk.color, cjk=_head_cjk if blk.role == 'heading' else None)
            else:
                ref = rest[0]
                if ref.kind == 'math':
                    p._p.append(parse_xml('<m:oMath%s>%s</m:oMath>' % (MNS, omml(ref.latex))))
                elif ref.kind == 'image':
                    try:
                        r = p.add_run()
                        r.add_picture(io.BytesIO(ref.png), width=Pt(ref.w))
                    except Exception:
                        pass
        if first_p is None:
            first_p = p
        first_inline = False
    return first_p

TW = 16.6

def render_docx(doc: Doc, path: str, progress=None):
    global TW
    TW = max(10.0, round((doc.width - doc.margin_l - doc.margin_r) / 28.35, 2))
    cm_l = doc.margin_l / 28.35
    cm_r = doc.margin_r / 28.35
    d = Document()
    s = d.sections[0]
    s.page_width = Cm(doc.width / 28.35)
    s.page_height = Cm(doc.height / 28.35)
    s.left_margin = Cm(max(1.2, min(cm_l, 4)))
    s.right_margin = Cm(max(1.2, min(cm_r, 4)))
    s.top_margin = Cm(max(1.2, min(doc.margin_t / 28.35, 4)))
    s.bottom_margin = Cm(max(1.2, min(doc.margin_b / 28.35, 4)))
    s.header_distance = Cm(0.6); s.footer_distance = Cm(1.1)

    nor = d.styles['Normal']
    nor.font.size = Pt(doc.body_size); nor.font.name = _body_latin
    nor.element.rPr.rFonts.set(qn('w:eastAsia'), _body_cjk)

    # ── 页眉/页脚：支持「左 / 中 / 右」三栏（原 PDF 的三栏结构）──
    def _hf_parts(blocks, n=3):
        out = []
        for b in blocks[:n]:
            out.append(' '.join(seg[1] for seg in parse_template(b.template, b.refs)
                                if seg[0] == 'text').strip())
        return out

    def _pagenum_like(t):
        return bool(re.match(r'^\s*\d{1,4}\s*(?:[/|·]\s*\d{1,4})?\s*$', t or ''))

    if doc.header:
        parts = _hf_parts(doc.header)
        hp = s.header.paragraphs[0]
        _clear(hp)
        hp = s.header.add_paragraph()
        hp.paragraph_format.tab_stops.add_tab_stop(Cm(TW / 2), WD_TAB_ALIGNMENT.CENTER)
        hp.paragraph_format.tab_stops.add_tab_stop(Cm(TW), WD_TAB_ALIGNMENT.RIGHT)
        if len(parts) >= 3:
            txt = '%s\t%s\t%s' % (parts[0], parts[1], parts[2])
        elif len(parts) == 2:
            txt = '%s\t%s' % (parts[0], parts[1])
        else:
            txt = (parts[0] if parts else '') + '\t'
        hp.add_run(_xml_safe(txt))
    # 页脚：三栏；中间那栏原文是页码就替换成 Word 原生 PAGE/NUMPAGES 域
    fp = s.footer.paragraphs[0]
    _clear(fp)
    fp = s.footer.add_paragraph()
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(TW / 2), WD_TAB_ALIGNMENT.CENTER)
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(TW), WD_TAB_ALIGNMENT.RIGHT)
    parts = _hf_parts(doc.footer)

    def _page_field(p):
        _field(p, 'PAGE')
        r2 = p.add_run(' / '); _set_run(r2, size=9)
        _field(p, 'NUMPAGES')

    if len(parts) >= 3:
        r = fp.add_run(_xml_safe(parts[0]) + '\t'); _set_run(r, size=9)
        if _pagenum_like(parts[1]):
            _page_field(fp)
        else:
            r = fp.add_run(_xml_safe(parts[1]) + '\t'); _set_run(r, size=9)
        r = fp.add_run('\t' + _xml_safe(parts[2])); _set_run(r, size=9)
    elif len(parts) == 2:
        r = fp.add_run(_xml_safe(parts[0]) + '\t'); _set_run(r, size=9)
        _page_field(fp)
        r = fp.add_run('\t' + _xml_safe(parts[1])); _set_run(r, size=9)
    else:
        r = fp.add_run(_xml_safe(parts[0] if parts else '') + '\t'); _set_run(r, size=9)
        _page_field(fp)

    def box(container, blk, nested=False):
        if blk.style == 'frame':
            bd = dict(top=(6, blk.border), bottom=(6, blk.border),
                      left=(6, blk.border), right=(6, blk.border))
        else:
            bd = dict(left=(18, blk.border))
        t = container.add_table(rows=1, cols=1)
        _props(t, fill=blk.fill)
        _borders(t, **bd)
        tw = t._tbl.tblPr.find(qn('w:tblW'))
        if nested:
            t.autofit = False; tw.set(qn('w:w'), '5000'); tw.set(qn('w:type'), 'pct')
        cell = t.cell(0, 0)
        _emit_cell(cell, blk.blocks)
        return t

    def _emit_cell(cell, blocks):
        first = True
        for b in blocks:
            if b.t == 'box':
                box(cell, b, nested=True)
            elif b.t == 'para':
                if first:
                    first = False
                    _clear(cell.paragraphs[0])
                _emit(cell, b)

    total = len(doc.blocks) or 1
    for i, blk in enumerate(doc.blocks):
        if progress:
            progress(0.9 + 0.1 * i / total, '排版中')
        if blk.t == 'pagebreak':
            # 原 PDF 的页边界：Word 里插真正的分页符。
            # （原先这里直接 continue，短页文档会被压缩合并成一页——实测 3 页压成 1 页）
            bp = d.add_paragraph()
            bp.paragraph_format.space_after = Pt(0)
            bp.paragraph_format.line_spacing = 1.0
            bp.add_run().add_break(WD_BREAK.PAGE)
            continue
        if blk.t == 'box':
            box(d, blk)
            sp = d.add_paragraph(); sp.paragraph_format.space_after = Pt(0)
            sp.paragraph_format.line_spacing = 1.0
        else:
            _emit(d, blk, space_after=4 if blk.role == 'heading' else 3)
    d.save(path)
    return path

# ───────────────────────── HTML ─────────────────────────
CSS = """
*{box-sizing:border-box}
body{margin:0;padding:22px 0;background:#eef0f3;color:#111;
 font-family:"Songti SC","STSong","宋体",serif;font-size:15px;line-height:1.75}
.page{width:21cm;min-height:29.7cm;background:#fff;margin:0 auto auto;padding:1.7cm 2.2cm 1.6cm;
 box-shadow:0 1px 6px rgba(0,0,0,.15)}
.runhead,.runfoot{display:flex;justify-content:space-between;font-size:11.5px;color:#333;
 border-bottom:.6px solid #000;padding-bottom:4px;margin-bottom:16px;gap:20px}
.runhead>span,.runfoot>span{white-space:nowrap}
/* 三栏页眉/页脚：中间那栏真正居中 */
.runhead>span:nth-child(2),.runfoot>span:nth-child(2){flex:1;text-align:center}
.runfoot{border-bottom:none;border-top:.6px solid #000;margin:20px 0 0;padding-top:4px}
p{margin:5px 0;text-align:justify}
h2,h3,h4{font-family:"Heiti SC","黑体",sans-serif;line-height:1.35;margin:14px 0 6px;color:#111}
h1{font-size:22px;text-align:center}
h2{font-size:17px}h3{font-size:15px}h4{font-size:13.5px}
.disp{margin:9px 0;text-align:center;overflow-x:auto}
ul.list{margin:6px 0;padding-left:1.4em;list-style:disc outside}
ul.list>li{margin:3px 0;text-align:justify}
.box ul.list{margin:4px 0}
math[display="block"]{font-size:1.06em}
math,.mmath{font-family:"Cambria Math","STIX Two Math",serif}
img.mmath{vertical-align:-2px}
.box{margin:9px 0;padding:9px 12px 10px;border-radius:1px}
.box .b-inner>.p-first{margin-top:0}
/* 原 PDF 的页边界：屏幕上看不出（高度 0），打印/导出 PDF 时强制分页 */
.pbreak{height:0;margin:0;padding:0;border:0;page-break-after:always;break-after:page}
@media print{body{background:#fff;padding:0}
 .page{width:auto!important;min-height:0!important;margin:0!important;padding:0!important;box-shadow:none!important}
 .pbreak{page-break-after:always;break-after:page}
 .box,table,tr{break-inside:avoid}
 h1,h2,h3,h4{break-after:avoid}
 /* 页眉/页脚：打印时隐藏，改由 render_pdf() 用 PyMuPDF 逐页盖章
    （CSS 的 fixed 定位在缩放打印下位置会错乱，实测页眉会跑到页面底部） */
 .runhead,.runfoot{display:none}
 @page{size:A4;margin:2.0cm 2.1cm 2.0cm}}
"""

def _seg_html(kind, rest, strip_lead_flag):
    """单个模板段 → HTML 片段。strip_lead_flag: 剥掉行首项目符号。"""
    if kind == 'text':
        txt, bold, italic = rest
        if strip_lead_flag:
            txt = strip_bullet(txt)[0]
        import html as H
        t = H.escape(_xml_safe(txt))
        if bold: t = '<b>%s</b>' % t
        if italic: t = '<i>%s</i>' % t
        return t
    ref = rest[0]
    if ref.kind == 'math':
        return mml_html(ref.latex)
    b64 = base64.b64encode(ref.png).decode()
    return ('<img class="mmath" src="data:image/png;base64,%s" '
            'style="height:%.2fpx">' % (b64, ref.h * 1.333))

def _html_segs(blk, strip_lead=False):
    out = []
    head = False
    for kind, *rest in parse_template(blk.template, blk.refs):
        flag = strip_lead and not head
        if kind == 'text':
            head = True
        out.append(_seg_html(kind, rest, flag))
    return ''.join(out)

def _big_img_html(ref, suffix_html=''):
    b64 = base64.b64encode(ref.png).decode()
    return ('<div class="disp"><img src="data:image/png;base64,%s" '
            'style="width:%.1fpx;max-width:96%%">%s</div>'
            % (b64, ref.w * 1.333, suffix_html))

def _html_rich(blk, strip_lead=False):
    """段落 → 若干 HTML 片段：超高行内图拆成独立居中大图（与 DOCX 一致）。"""
    chunks = []
    stripped = False
    for typ, *rest in _split_big(blk):
        if typ == 'big':
            ref, suffix = rest[0], rest[1]
            sfx = ''.join(_seg_html(k, r, False) for k, *r in suffix)
            chunks.append(('big', _big_img_html(ref, sfx)))
            continue
        payload = rest[0]
        out = []
        for k, *r in payload:
            flag = strip_lead and not stripped and k == 'text'
            if flag:
                stripped = True
            out.append(_seg_html(k, r, flag))
        chunks.append(('inline', ''.join(out)))
    return chunks


def _html_para(blk, doc=None):
    """单个非列表块 → HTML。"""
    if blk.display:
        return '<div class="disp">%s</div>' % _html_segs(blk)
    if blk.role == 'heading':
        lv = 2 if blk.size > (doc.body_size if doc else 11) + 6 else 3
        return '<h%d style="color:#%s">%s</h%d>' % (lv, blk.color, _html_segs(blk), lv)
    css_p = 'style="color:#%s"' % blk.color if blk.color and blk.color != '000000' else ''
    if blk.role == 'list':
        body = ''.join(h for _, h in _html_rich(blk, strip_lead=True))
        return '<li %s>%s</li>' % (css_p, body)
    out = []
    for typ, h in _html_rich(blk):
        out.append(h if typ == 'big' else '<p %s>%s</p>' % (css_p, h))
    return ''.join(out)


def _html_blocks(blocks, doc=None):
    """一串块 → HTML：连续的列表项合并成真正的 <ul>，而不是把 • 挤在同一行。"""
    out, i = [], 0
    while i < len(blocks):
        blk = blocks[i]
        if blk.t == 'pagebreak':
            # 保留原 PDF 的页边界：显式分页（Chrome 打印 PDF 时生效）。
            # 原先直接跳过 → HTML/PDF 全篇连排，短页文档会被压成一页。
            out.append('<div class="pbreak"></div>')
            i += 1
            continue
        if blk.t == 'box':
            out.append(_html_box(blk, doc))
            i += 1
            continue
        if blk.t == 'para' and blk.role == 'list':
            items = []
            while i < len(blocks) and blocks[i].t == 'para' and blocks[i].role == 'list':
                items.append(''.join(h for _, h in _html_rich(blocks[i], strip_lead=True)))
                i += 1
            out.append('<ul class="list">%s</ul>'
                       % ''.join('<li>%s</li>' % s for s in items))
            continue
        out.append(_html_para(blk, doc))
        i += 1
    return ''.join(out)

def render_html(doc: Doc, path: str, progress=None, zoom: float = None):
    """zoom<1 时只作用于「打印/导出 PDF」：等比缩小内容，
    让原 PDF 的每一页尽量落在一张 A4 上（页数贴近原文件、减少半空白页）。
    屏幕浏览（尤其是 html 产物）不受影响。"""
    css = CSS
    if zoom and zoom < 0.999:
        css = css.replace('@media print{body{background:#fff;padding:0}',
                          '@media print{body{background:#fff;padding:0;zoom:%.3f}' % zoom)
    W = doc.width - doc.margin_l - doc.margin_r
    body = []
    if doc.header:
        hs = [_html_segs(b) for b in doc.header][:3]
        while len(hs) < 2:
            hs.append('')
        body.append('<div class="runhead">%s</div>'
                    % ''.join('<span>%s</span>' % h for h in hs))
    body.append(_html_blocks(doc.blocks, doc))
    if doc.footer or doc.has_pagenum:
        fs = [_html_segs(b) for b in doc.footer][:3]
        fs = [f for f in fs if f.strip()]
        while len(fs) < 2:
            fs.append('')
        if doc.has_pagenum:
            # 页码放中间（与原 PDF 的「左 / 中 / 右」三栏一致）。
            # 注意：HTML/PDF 是静态页面，拿不到"当前第几页"，硬写会每页都一样；
            # 真实页码只在 DOCX 里用 Word 的 PAGE/NUMPAGES 域给出。
            pn = ''
            if len(fs) >= 3:
                fs = [fs[0], pn, fs[2]]
            elif len(fs) == 2:
                fs = [fs[0], pn, fs[1]]
            else:
                fs = fs + [pn]
        body.append('<div class="runfoot">%s</div>'
                    % ''.join('<span>%s</span>' % f for f in fs))
    html = ('<!DOCTYPE html><html lang="zh-CN"><head><meta charset="utf-8">'
            '<title>译文</title><style>%s</style></head><body><div class="page" '
            'style="width:%.1fcm;padding-left:%.2fcm;padding-right:%.2fcm">%s</div></body></html>'
            % (css, doc.width / 28.35, doc.margin_l / 28.35, doc.margin_r / 28.35,
               '\n'.join(body)))
    open(path, 'w', encoding='utf-8').write(html)
    return path

# ───────────────────────── PDF ─────────────────────────
_CHROME = ('/Applications/Google Chrome.app/Contents/MacOS/Google Chrome',
           '/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge',
           '/Applications/Chromium.app/Contents/MacOS/Chromium',
           'google-chrome', 'chromium', 'chromium-browser')

def find_browser():
    import shutil
    for p in _CHROME:
        if os.path.exists(p):
            return p
        q = shutil.which(p)
        if q:
            return q
    return None

def render_pdf_story(html_path: str, out: str, css_extra: str = ''):
    """HTML → PDF（不依赖浏览器）：用 PyMuPDF 的 Story 引擎排版，文字可选中、矢量输出。

    服务器环境通常没有 Chrome，这是云端部署时的备选路径。
    """
    import pymupdf
    html = open(html_path, encoding='utf-8').read()
    # Story 不支持 flex / box-shadow 等现代属性，做一次等价降级
    html = html.replace('display:flex;justify-content:space-between',
                        'display:block;text-align:left')
    html = html.replace('display:flex', 'display:block')
    # 该引擎不认布局视口右边距、对 margin/padding 支持不完整，只认元素宽度：
    # 页面盒固定 500pt、左移 47pt；行内公式图与字号微缩，抵消 fallback 字体的度量差
    html = re.sub(r'(<div class="page" style=")[^"]*(")',
                  r'\1width:500pt;margin-left:47pt\2', html)
    html = re.sub(r'width="(\d+(?:\.\d+)?)"',
                  lambda m: 'width="%.0f"' % (float(m.group(1)) * 0.90), html)
    html = html.replace('font-size:15px;line-height:1.75',
                        'font-size:12px;line-height:1.75')
    story = pymupdf.Story(html, user_css=css_extra)
    writer = pymupdf.DocumentWriter(out)
    med = pymupdf.paper_rect('a4')
    guard = 0
    more = True
    while more and guard < 800:
        dev = writer.begin_page(med)
        more, _ = story.place(med)
        story.draw(dev)
        writer.end_page()
        guard += 1
    writer.close()
    return out


def render_pdf(doc: Doc, path: str, html_path: str = None, timeout=180):
    """HTML → PDF：优先用 Chrome/Edge 无头打印，失败则退到 PyMuPDF Story。"""
    import subprocess, tempfile
    browser = find_browser()
    if not browser:
        if html_path:
            render_pdf_story(html_path, path)
            return path
        raise RuntimeError('未找到 Chrome/Edge 浏览器，无法生成 PDF（可先只选 Word / HTML）')
    tmp = html_path
    if not tmp:
        fd, tmp = tempfile.mkstemp(suffix='.html')
        os.close(fd)
        render_html(doc, tmp)
    out = os.path.abspath(path)
    cmd = [browser, '--headless=new', '--disable-gpu', '--no-sandbox',
           '--no-pdf-header-footer', '--run-all-compositor-stages-before-draw',
           '--virtual-time-budget=15000',
           '--print-to-pdf=' + out,
           'file://' + os.path.abspath(tmp)]
    env = dict(os.environ)
    env.pop('HTTP_PROXY', None); env.pop('HTTPS_PROXY', None)
    env.pop('http_proxy', None); env.pop('https_proxy', None)
    try:
        r = subprocess.run(cmd, capture_output=True, timeout=timeout, env=env)
    except Exception:
        r = None
    if not os.path.exists(out):
        # 浏览器打印失败（服务器环境常见）→ 退到内置排版引擎
        render_pdf_story(tmp, out)
    try:
        stamp_running(doc, out)
    except Exception:
        pass
    return out


def _hf_text(blocks, n=3):
    out = []
    for b in blocks[:n]:
        out.append(' '.join(seg[1] for seg in parse_template(b.template, b.refs)
                            if seg[0] == 'text').strip())
    return out


def stamp_running(doc: Doc, pdf_path: str):
    """给 PDF 每一页盖页眉/页脚（左/中/右三栏 + 真实页码）。

    为什么不在 HTML 里做：CSS 的 position:fixed 在 Chrome 缩放打印下位置会错乱
    （实测页眉跑到页面底部），而静态 HTML 又拿不到"当前第几页"。生成 PDF 后
    逐页盖章最稳，还能给出真实页码 "n / N"。
    """
    import pymupdf
    d = pymupdf.open(pdf_path)
    n = len(d)
    hd, ft = _hf_text(doc.header), _hf_text(doc.footer)
    if not hd and not ft and not doc.has_pagenum:
        return pdf_path
    FONT = 'china-ss'          # PyMuPDF 内置简体宋体，中英混排都能画
    FS = 9
    for i, page in enumerate(d):
        W, H = page.rect.width, page.rect.height
        ml = max(28.0, doc.margin_l * 0.85)
        mr = W - max(28.0, doc.margin_r * 0.85)
        mid = (ml + mr) / 2

        def put(txt, x, y, align='left', fs=FS):
            if not txt:
                return
            w = pymupdf.get_text_length(txt, fontname=FONT, fontsize=fs)
            if align == 'right':
                x -= w
            elif align == 'center':
                x -= w / 2
            page.insert_text((x, y), txt, fontname=FONT, fontsize=fs, color=(0.2, 0.2, 0.2))

        def fit(txt, avail, fs=FS):
            """左侧文字可用宽度不够时按比例缩小字号（宋体偏宽，长页眉会压到右栏）。"""
            if not txt:
                return fs
            w = pymupdf.get_text_length(txt, fontname=FONT, fontsize=fs)
            return max(5.5, min(fs, fs * avail / w)) if w > avail else fs

        # 页眉：左 / 右（原档多为两栏），下方一条细线
        if hd:
            y = 34
            if len(hd) >= 3:
                wm = pymupdf.get_text_length(hd[1], fontname=FONT, fontsize=FS)
                wr = pymupdf.get_text_length(hd[2], fontname=FONT, fontsize=FS)
                gap = (mr - ml - wr - wm) / 2
                fs_m = fit(hd[1], max(30, gap), FS)
                fs_l = fit(hd[0], max(40, (mr - ml) - wr - wm - 8), FS)
                put(hd[0], ml, y, fs=fs_l)
                put(hd[1], mid, y, 'center', fs_m)
                put(hd[2], mr, y, 'right')
            elif len(hd) == 2:
                fs_l = fit(hd[0], max(40, (mr - ml) -
                                      pymupdf.get_text_length(hd[1], fontname=FONT,
                                                              fontsize=FS) - 8), FS)
                put(hd[0], ml, y, fs=fs_l)
                put(hd[1], mr, y, 'right')
            page.draw_line(pymupdf.Point(ml, y + 5), pymupdf.Point(mr, y + 5),
                           color=(0, 0, 0), width=0.5)
        # 页脚：左 / 页码 / 右，上方一条细线
        if ft or doc.has_pagenum:
            y = H - 30
            page.draw_line(pymupdf.Point(ml, y - 12), pymupdf.Point(mr, y - 12),
                           color=(0, 0, 0), width=0.5)
            pn = '%d / %d' % (i + 1, n) if doc.has_pagenum else ''
            wpn = pymupdf.get_text_length(pn, fontname=FONT, fontsize=FS) if pn else 0
            if len(ft) >= 3:
                fs_l = fit(ft[0], max(40, (mr - ml) - wpn -
                                      pymupdf.get_text_length(ft[2], fontname=FONT,
                                                              fontsize=FS) - 16), FS)
                put(ft[0], ml, y, fs=fs_l)
                put(ft[2], mr, y, 'right')
            elif len(ft) == 2:
                fs_l = fit(ft[0], max(40, (mr - ml) - wpn -
                                      pymupdf.get_text_length(ft[1], fontname=FONT,
                                                              fontsize=FS) - 16), FS)
                put(ft[0], ml, y, fs=fs_l)
                put(ft[1], mr, y, 'right')
            elif len(ft) == 1:
                fs_l = fit(ft[0], max(40, (mr - ml) - wpn - 16), FS)
                put(ft[0], ml, y, fs=fs_l)
            if doc.has_pagenum:
                put(pn, mid, y, 'center')
    tmp = pdf_path + '.stamped'
    d.save(tmp, garbage=0, deflate=True)
    d.close()
    os.replace(tmp, pdf_path)
    return pdf_path

def _html_box(blk, doc=None):
    """盒子（tcolorbox）→ HTML，盒内同样支持列表分组。"""
    css = ('border:1px solid #%s;background:#%s;' % (blk.border, blk.fill)
           if blk.style == 'frame'
           else 'border-left:3px solid #%s;background:#%s;' % (blk.border, blk.fill))
    inner = []
    for b in blk.blocks:
        if b.t == 'box':
            inner.append(_html_box(b, doc))
        elif b.t == 'para':
            if b.role == 'heading':
                st = 'font-weight:700;color:#%s;font-size:%.1fpx' % (
                    b.color, max(14, b.size * 1.35))
            elif b.color and b.color != '000000':
                st = 'color:#%s' % b.color
            else:
                st = ''
            if b.display:
                inner.append('<div class="disp">%s</div>' % _html_segs(b))
            elif b.role == 'list':
                inner.append('<ul class="list"><li style="%s">%s</li></ul>'
                             % (st, ''.join(h for _, h in _html_rich(b, strip_lead=True))))
            else:
                for typ, h in _html_rich(b):
                    inner.append(h if typ == 'big' else '<p style="%s">%s</p>' % (st, h))
    return '<div class="box" style="%s">%s</div>' % (css, ''.join(inner))
